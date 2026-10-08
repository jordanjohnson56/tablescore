"""Build data/seed.json from the rubric workbook and BGG collection export.

Usage: python3 scripts/build_seed.py RUBRIC.xlsx collection.csv > data/seed.json
"""
import csv, json, sys
import openpyxl

CRITERIA = ["desire", "table", "depth", "interaction", "replay", "art", "theme"]


def num(v, cast=float):
    try:
        return cast(v) if v not in (None, "") else None
    except ValueError:
        return None


def main(xlsx, csv_path):
    wb = openpyxl.load_workbook(xlsx)
    w = wb["Weights"]
    weights = {k.lower(): w.cell(row=r, column=3).value
               for r, k in zip(range(4, 11), ["Desire", "Table", "Depth", "Interaction", "Replay", "Art", "Theme"])}
    stretch = w["C14"].value

    bgg = {}
    with open(csv_path, newline="", encoding="utf-8") as f:
        for r in csv.DictReader(f):
            bgg[int(r["objectid"])] = r

    games = []
    for row in wb["Scores"].iter_rows(min_row=2, values_only=True):
        name, bgg_id, typ, status, plays, old, avg, calib = row[:8]
        if not name:
            continue
        scores = {c: num(v) for c, v in zip(CRITERIA, row[8:15])}
        b = bgg.get(bgg_id, {})
        games.append({
            "id": f"bgg-{bgg_id}" if bgg_id else f"g-{len(games)}",
            "name": name,
            "bggId": bgg_id,
            "type": typ or "",
            "status": status or "",
            "plays": plays or 0,
            "oldRating": num(old),
            "bggAvg": num(avg),
            "calibration": calib == "Yes",
            "scores": scores,
            "avoidTheme": row[15] == "Yes",
            "notes": row[21] or "",
            "year": num(b.get("yearpublished"), int),
            "minPlayers": num(b.get("minplayers"), int),
            "maxPlayers": num(b.get("maxplayers"), int),
            "bestPlayers": b.get("bggbestplayers") or "",
            "playTime": num(b.get("playingtime"), int),
            "bggWeight": round(num(b.get("avgweight")) or 0, 2) or None,
        })

    json.dump({"version": 1, "settings": {"weights": weights, "stretch": stretch}, "games": games},
              sys.stdout, ensure_ascii=False, indent=1)


if __name__ == "__main__":
    main(*sys.argv[1:3])
